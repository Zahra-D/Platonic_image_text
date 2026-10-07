"""Append metric and back-translation explanation slides to the CLEVR deck."""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


DECK = Path("/home/zd25e122/Omni/presentation/CLEVR_stage0_group_presentation.pptx")
BACKUP = DECK.with_name("CLEVR_stage0_group_presentation.before_metric_translation_slides.pptx")

NAVY = RGBColor(0x0B, 0x13, 0x2B)
PANEL = RGBColor(0x11, 0x1D, 0x3A)
MUTED = RGBColor(0xAA, 0xB6, 0xD3)
WHITE = RGBColor(0xF4, 0xF7, 0xFF)
GRID = RGBColor(0x30, 0x42, 0x63)
CYAN = RGBColor(0x45, 0xD7, 0xE8)
GREEN = RGBColor(0x72, 0xE6, 0xA5)
YELLOW = RGBColor(0xF3, 0xC9, 0x69)
PURPLE = RGBColor(0xA7, 0x8B, 0xFA)
RED = RGBColor(0xFF, 0x7B, 0x72)


def box(slide, x, y, w, h, fill=PANEL, line=GRID, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.2)
    return shape


def text(slide, value, x, y, w, h, size=12, color=WHITE, bold=False,
         align=PP_ALIGN.LEFT, margin=0.03, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def rich_lines(slide, lines, x, y, w, h, size=11.2, spacing=4):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.08)
    frame.margin_top = frame.margin_bottom = Inches(0.05)
    for index, (lead, rest, color) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(spacing)
        paragraph.level = 0
        r1 = paragraph.add_run()
        r1.text = lead
        r1.font.name = "Aptos"
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = color
        r2 = paragraph.add_run()
        r2.text = rest
        r2.font.name = "Aptos"
        r2.font.size = Pt(size)
        r2.font.color.rgb = WHITE
    return shape


def header(slide, eyebrow, title_value, number):
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    text(slide, eyebrow, 0.62, 0.31, 9.8, 0.28, 9.5, CYAN, True)
    text(slide, title_value, 0.62, 0.67, 11.65, 0.58, 27, WHITE, True)
    text(slide, f"{number:02d}", 12.25, 0.36, 0.45, 0.28, 10, MUTED, True, PP_ALIGN.RIGHT)
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.38), Inches(12.05), Inches(0.02))
    divider.fill.solid(); divider.fill.fore_color.rgb = GRID; divider.line.fill.background()


def footer(slide, value):
    text(slide, value, 0.62, 7.13, 11.8, 0.20, 8.5, MUTED)


def add_metric_slide(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "EVALUATION PROTOCOL", "Text → image: exactly what each loss scores", number)
    text(slide, "Target = Iᵢ · caption supplies context · lower token NLL is better", 0.72, 1.48, 11.7, 0.30, 12.0, MUTED)

    rows = [
        ("METRIC", "MODEL INPUT", "GROUND TRUTH SCORED", CYAN),
        ("MATCHED", "clean Tᵢ  +  masked Iᵢ", "original tokens of Iᵢ", GREEN),
        ("SHUFFLED", "clean Tᵢ₊₁  +  the same masked Iᵢ", "original tokens of Iᵢ", RED),
        ("NULL", "masked Iᵢ only", "original tokens of Iᵢ", MUTED),
    ]
    x, y = 0.72, 1.92
    widths = [1.55, 6.25, 4.00]
    heights = [0.58, 0.90, 0.90, 0.90]
    for row_index, row in enumerate(rows):
        cx = x
        for col_index, width in enumerate(widths):
            fill = RGBColor(0x18, 0x29, 0x4D) if row_index == 0 else PANEL
            box(slide, cx, y, width, heights[row_index], fill, GRID, radius=False)
            color = row[3] if col_index == 0 and row_index else (CYAN if row_index == 0 else WHITE)
            text(
                slide, row[col_index], cx + 0.12, y + 0.10, width - 0.24, heights[row_index] - 0.18,
                10.3 if row_index == 0 else 12.0, color, row_index == 0 or col_index == 0,
                valign=MSO_ANCHOR.MIDDLE,
            )
            cx += width
        y += heights[row_index]

    box(slide, 0.72, 5.55, 11.80, 1.00, PANEL, YELLOW)
    rich_lines(slide, [
        ("Shuffling rule — ", "caption comes from the next batch row, with wraparound.", YELLOW),
        ("Controlled comparison — ", "the image target and its mask never change; only caption identity changes.", CYAN),
    ], 0.94, 5.76, 11.35, 0.62, 11.5, 4)
    footer(slide, "Implementation: alignment_evaluation.py · shuffled source row = (i + 1) mod batch size")


def add_translation_route_slide(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "UNPAIRED TRANSLATION", "How a pseudo-pair is generated and routed", number)
    text(slide, "Real carriers are deranged: Tᵢ and Iⱼ are deliberately unrelated (i ≠ j).", 0.72, 1.48, 11.6, 0.30, 12, MUTED)

    stages = [
        ("1 · SOURCE", "real Tᵢ or Iⱼ\nsource-private OFF", CYAN),
        ("2 · TARGET", "all semantic tokens masked\ntarget private ON", MUTED),
        ("3 · GENERATE", "12 greedy steps\nconfidence reveal", YELLOW),
        ("4 · PSEUDO-PAIR", "(Tᵢ, Îᵢ) or (T̂ⱼ, Iⱼ)\ndiscrete tokens", PURPLE),
    ]
    x_positions = [0.72, 3.72, 6.72, 9.72]
    for (title_value, body, accent), x in zip(stages, x_positions):
        box(slide, x, 1.96, 2.65, 1.22, PANEL, accent)
        text(slide, title_value, x + 0.17, 2.12, 2.32, 0.24, 11.5, accent, True)
        text(slide, body, x + 0.17, 2.45, 2.32, 0.55, 10.2, WHITE)

    box(slide, 0.72, 3.55, 5.75, 2.65, PANEL, GREEN)
    text(slide, "TEXT → IMAGE", 0.98, 3.80, 5.20, 0.30, 15, GREEN, True)
    rich_lines(slide, [
        ("Condition Tᵢ — ", "shared-only route (private branch disabled).", CYAN),
        ("Target image — ", "masked image positions use shared + image-private.", YELLOW),
        ("Reveal — ", "temperature 0 argmax; highest-confidence tokens first.", PURPLE),
        ("Important — ", "Iⱼ supplies only target positions; its semantic tokens are masked.", RED),
    ], 0.98, 4.25, 5.18, 1.65, 10.7, 5)

    box(slide, 6.75, 3.55, 5.77, 2.65, PANEL, PURPLE)
    text(slide, "IMAGE → TEXT", 7.01, 3.80, 5.20, 0.30, 15, PURPLE, True)
    rich_lines(slide, [
        ("Condition Iⱼ — ", "shared-only route (private branch disabled).", CYAN),
        ("Target text — ", "masked lexical positions use shared + text-private.", YELLOW),
        ("Special tokens — ", "modality/BOS/EOS remain as the target scaffold.", MUTED),
        ("Filter — ", "pseudo-row retained only when mean confidence ≥ 0.10.", GREEN),
    ], 7.01, 4.25, 5.18, 1.65, 10.7, 5)
    footer(slide, "Translation config: pseudo-batch 8 · every 10 microsteps · 12 steps · confidence order · temperature 0")


def add_translation_loss_slide(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "LOSS + BACKPROP", "What translation optimizes—and where gradients flow", number)

    box(slide, 0.72, 1.63, 5.75, 3.72, PANEL, GREEN)
    text(slide, "A · CYCLE RECONSTRUCTION", 0.98, 1.88, 5.18, 0.32, 15, GREEN, True)
    text(slide, "Lcycle = CE(original source tokens, reconstructed source logits)", 0.98, 2.35, 5.17, 0.55, 13.2, WHITE, True)
    rich_lines(slide, [
        ("Input — ", "generated pseudo-target is clean context; original source is randomly masked.", CYAN),
        ("Routing — ", "pseudo-target shared-only; reconstructed source shared + its private branch.", YELLOW),
        ("Ground truth — ", "the original real source tokens, never the unrelated carrier.", PURPLE),
        ("Gradient — ", "cycle CE updates the reconstruction forward; discrete generation is detached.", RED),
    ], 0.98, 3.05, 5.18, 1.88, 10.7, 5)

    box(slide, 6.75, 1.63, 5.77, 3.72, PANEL, PURPLE)
    text(slide, "B · SHARED InfoNCE", 7.01, 1.88, 5.20, 0.32, 15, PURPLE, True)
    text(slide, "Lalign = CE(sim(hpseudo, hsource) / τ, matching row)", 7.01, 2.35, 5.17, 0.55, 13.2, WHITE, True)
    rich_lines(slide, [
        ("Representations — ", "both source and pseudo-target are encoded shared-only.", CYAN),
        ("Positive — ", "same pseudo-pair row; other accepted rows are negatives; τ = 0.07.", YELLOW),
        ("Teacher — ", "real-source representation uses no_grad and is detached.", MUTED),
        ("Gradient — ", "flows through pseudo-target shared path; no private branch receives Lalign.", GREEN),
    ], 7.01, 3.05, 5.18, 1.88, 10.7, 5)

    box(slide, 0.72, 5.65, 11.80, 0.88, PANEL, CYAN)
    text(slide, "LBT,direction = 0.10 Lcycle + 0.05 Lalign", 0.95, 5.83, 4.35, 0.30, 13.5, CYAN, True)
    text(slide, "Average valid directions × warmup min(1, (step+1)/1000)", 5.20, 5.83, 7.05, 0.30, 11.3, WHITE)
    footer(slide, "No gradient passes through argmax/confidence generation · normal marginal task loss and DANN are optimized in the same run")


def main():
    if not DECK.is_file():
        raise FileNotFoundError(DECK)
    prs = Presentation(DECK)
    existing_text = "\n".join(
        shape.text for slide in prs.slides for shape in slide.shapes
        if hasattr(shape, "text")
    )
    marker = "Text → image: exactly what each loss scores"
    if marker in existing_text:
        raise RuntimeError("Metric/translation slides already exist; refusing to duplicate them")
    if not BACKUP.exists():
        shutil.copy2(DECK, BACKUP)
    start = len(prs.slides) + 1
    add_metric_slide(prs, start)
    add_translation_route_slide(prs, start + 1)
    add_translation_loss_slide(prs, start + 2)
    prs.save(DECK)
    print(f"saved {DECK} with {len(prs.slides)} slides; backup={BACKUP}")


if __name__ == "__main__":
    main()
